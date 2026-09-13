# -*- coding: utf-8 -*-
"""خروجی اکسل — کارپوشهٔ قالب‌بندی‌شده، نه یک فایل CSV با پسوند xlsx.

آنچه خروجی را «گزارش» می‌کند و نه «داده»:

* برگهٔ «خلاصه» با شاخص‌های کلیدی و سرصفحهٔ فراداده.
* برگهٔ «کیفیت داده» با چهار زیرامتیاز و وزن‌ها.
* برگهٔ «پروفایل ستون‌ها» و «مسائل و اصلاحات» برای بازرسی کامل.
* برگهٔ «داده پاک‌شده» با سرصفحهٔ ثابت، فیلتر خودکار، عرض ستون مناسب و
  قالب عددی — کاربر بتواند بی‌درنگ روی آن کار کند.

جهت برگه‌ها راست‌به‌چپ است (``sheet_view.rightToLeft``). بدون آن، اکسل جدول
فارسی را چپ‌به‌راست نشان می‌دهد و ستون تاریخ در سمت اشتباه می‌نشیند.
"""
from __future__ import annotations

import io
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analytics import Analysis, movers
from .cleaning import CleanResult
from .config import (OUTPUT_BASENAMES, SHEET_COLUMNS, SHEET_DATA,
                     SHEET_ISSUES, SHEET_QUALITY, SHEET_SUMMARY)
from .report import Report

#: سقف سطرهایی که در برگهٔ داده نوشته می‌شود — اکسل با میلیون سطر عملاً باز
#: نمی‌شود و فایل خروجی هم غیرقابل استفاده می‌شود.
MAX_EXPORT_ROWS = 60_000
#: سقف پهنای ستون در اکسل.
MAX_COLUMN_WIDTH = 46

THIN = Side(style="thin", color="D8DEE8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color.lstrip("#"))


def _safe_sheet_title(title: str) -> str:
    """نام برگه: اکسل حداکثر ۳۱ نویسه می‌پذیرد و ``[]:*?/\\`` را رد می‌کند."""
    cleaned = "".join("-" if char in "[]:*?/\\" else char for char in str(title))
    return cleaned[:31] or "برگه"


def _write_header(sheet, row: int, headers: list[str], theme: dict) -> None:
    """نوشتن سرصفحهٔ یک جدول با رنگ تم."""
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index, value=str(title))
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = _fill(theme["primary"])
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def _autosize(sheet, headers: list[str], rows: list[list]) -> None:
    """عرض ستون از بلندترین محتوا — با سقف."""
    for index, title in enumerate(headers, start=1):
        longest = max([len(str(title))] +
                      [len(str(row[index - 1])) for row in rows
                       if len(row) >= index and row[index - 1] is not None] or [8])
        sheet.column_dimensions[get_column_letter(index)].width = min(
            MAX_COLUMN_WIDTH, max(9, longest + 3))


def _table(sheet, row: int, headers: list[str], rows: list[list],
           theme: dict, *, add_filter: bool = True) -> int:
    """درج یک جدول کامل و برگرداندن سطر بعدی آزاد."""
    _write_header(sheet, row, headers, theme)
    for offset, values in enumerate(rows, start=1):
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row + offset, column=index)
            cell.value = _coerce(value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if offset % 2 == 0:
                cell.fill = _fill(theme["soft"])
    if add_filter and rows:
        sheet.auto_filter.ref = (f"A{row}:"
                                 f"{get_column_letter(len(headers))}{row + len(rows)}")
    _autosize(sheet, headers, rows)
    return row + len(rows) + 2


def _coerce(value):
    """مقدار سلول: انواع pandas و numpy به انواع بومی اکسل تبدیل می‌شوند."""
    if value is None:
        return None
    if value is pd.NaT:
        return None
    if isinstance(value, float) and value != value:
        return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            return str(value)
    return value


def _cover(sheet, report: Report, theme: dict) -> int:
    """سرصفحهٔ گزارش: عنوان، تم و فراداده."""
    sheet.merge_cells("A1:D1")
    title = sheet["A1"]
    title.value = report.title
    title.font = Font(bold=True, size=18, color="FFFFFF")
    title.fill = _fill(theme["primary"])
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:D2")
    subtitle = sheet["A2"]
    subtitle.value = report.subtitle
    subtitle.font = Font(size=11, color=theme["muted"].lstrip("#"))
    subtitle.alignment = Alignment(horizontal="center")

    row = 4
    for label, value in report.meta_lines:
        sheet.cell(row=row, column=1, value=str(label)).font = Font(bold=True)
        cell = sheet.cell(row=row, column=2, value=str(value))
        cell.alignment = Alignment(horizontal="right")
        for column in (1, 2):
            sheet.cell(row=row, column=column).border = BORDER
        row += 1
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 44
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 20
    return row + 1


def build_workbook(report: Report, *, analysis: Analysis, clean: CleanResult | None,
                   frame: pd.DataFrame | None) -> bytes:
    """ساخت کارپوشهٔ اکسل گزارش."""
    theme = report.theme
    book = Workbook()

    #: ---------------------------------------------------------- خلاصه
    summary = book.active
    summary.title = _safe_sheet_title(SHEET_SUMMARY)
    summary.sheet_view.rightToLeft = True
    row = _cover(summary, report, theme)

    kpi_items = []
    kpi_section = report.section("kpis")
    for block in (kpi_section.blocks if kpi_section else []):
        if block.kind == "kpis":
            kpi_items = block.data.get("entries", [])
    if kpi_items:
        summary.cell(row=row, column=1, value="شاخص‌های کلیدی").font = Font(
            bold=True, size=13, color=theme["primary"].lstrip("#"))
        row += 1
        row = _table(summary, row,
                     ["شاخص", "مقدار", "تغییر", "توضیح"],
                     [[item.get("label", ""), item.get("value", ""),
                       (f"{item.get('change', '')} "
                        f"{'افزایش' if item.get('direction') == 'up' else 'کاهش'}"
                        if item.get("change") else "—"),
                       item.get("hint", "")]
                      for item in kpi_items], theme, add_filter=False)

    summary_block = report.section("summary")
    for block in (summary_block.blocks if summary_block else []):
        if block.kind == "paragraph":
            summary.cell(row=row, column=1, value="خلاصهٔ مدیریتی").font = Font(
                bold=True, size=13, color=theme["primary"].lstrip("#"))
            row += 1
            summary.merge_cells(start_row=row, start_column=1, end_row=row + 3,
                                end_column=4)
            cell = summary.cell(row=row, column=1, value=block.data.get("text", ""))
            cell.alignment = Alignment(horizontal="right", vertical="top",
                                       wrap_text=True)
            summary.row_dimensions[row].height = 72
            row += 5

    #: ---------------------------------------------------------- کیفیت
    quality_sheet = book.create_sheet(_safe_sheet_title(SHEET_QUALITY))
    quality_sheet.sheet_view.rightToLeft = True
    quality_section = report.section("quality")
    row = 2
    for block in (quality_section.blocks if quality_section else []):
        if block.kind == "table":
            row = _table(quality_sheet, row, block.data.get("columns", []),
                         block.data.get("rows", []), theme)
        elif block.kind == "kpis":
            items = block.data.get("entries", [])
            row = _table(quality_sheet, row, ["شاخص", "مقدار", "توضیح"],
                         [[item.get("label", ""), item.get("value", ""),
                           item.get("hint", "")] for item in items], theme)
        elif block.kind == "callout":
            quality_sheet.cell(row=row, column=1,
                               value=block.data.get("title", "")).font = Font(bold=True)
            row += 1
            quality_sheet.merge_cells(start_row=row, start_column=1,
                                      end_row=row + 3, end_column=4)
            cell = quality_sheet.cell(row=row, column=1,
                                      value=block.data.get("body", ""))
            cell.alignment = Alignment(horizontal="right", vertical="top",
                                       wrap_text=True)
            quality_sheet.row_dimensions[row].height = 66
            row += 5

    #: ---------------------------------------------------------- پروفایل ستون‌ها
    columns_sheet = book.create_sheet(_safe_sheet_title(SHEET_COLUMNS))
    columns_sheet.sheet_view.rightToLeft = True
    appendix = report.section("appendix")
    row = 2
    for block in (appendix.blocks if appendix else []):
        if block.kind == "table":
            row = _table(columns_sheet, row, block.data.get("columns", []),
                         block.data.get("rows", []), theme)

    #: ---------------------------------------------------------- مسائل و اصلاحات
    issues_sheet = book.create_sheet(_safe_sheet_title(SHEET_ISSUES))
    issues_sheet.sheet_view.rightToLeft = True
    row = 2
    if clean and clean.findings:
        issues_sheet.cell(row=row, column=1,
                          value="مسائل پیداشده").font = Font(bold=True, size=13)
        row += 1
        row = _table(issues_sheet, row, ["مسئله", "شرح", "تعداد"],
                     [[item.title, item.detail, item.affected]
                      for item in clean.findings], theme)
    if clean and clean.fixes:
        issues_sheet.cell(row=row, column=1,
                          value="اصلاحات اعمال‌شده").font = Font(bold=True, size=13)
        row += 1
        _table(issues_sheet, row, ["اصلاح", "شرح", "تعداد"],
               [[item.title, item.detail, item.affected] for item in clean.fixes],
               theme)
    if not (clean and (clean.findings or clean.fixes)):
        issues_sheet.cell(row=row, column=1, value="مسئله یا اصلاحی ثبت نشد.")

    #: ---------------------------------------------------------- داده پاک‌شده
    if frame is not None and len(frame.columns):
        data_sheet = book.create_sheet(_safe_sheet_title(SHEET_DATA))
        data_sheet.sheet_view.rightToLeft = True
        headers = [str(name) for name in frame.columns]
        _write_header(data_sheet, 1, headers, theme)
        trimmed = frame.head(MAX_EXPORT_ROWS)
        for row_index, (_index, values) in enumerate(trimmed.iterrows(), start=2):
            for column_index, name in enumerate(frame.columns, start=1):
                cell = data_sheet.cell(row=row_index, column=column_index)
                cell.value = _coerce(values[name])
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if isinstance(cell.value, float) \
                        else "#,##0"
        data_sheet.freeze_panes = "A2"
        if len(trimmed):
            data_sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}{len(trimmed) + 1}")
        for index, name in enumerate(frame.columns, start=1):
            widths = [len(str(name))]
            if name in trimmed.columns:
                widths.extend(len(value)
                              for value in trimmed[name].astype(str).head(400))
            data_sheet.column_dimensions[get_column_letter(index)].width = min(
                MAX_COLUMN_WIDTH, max(10, max(widths) + 2))
        if len(frame) > MAX_EXPORT_ROWS:
            note_row = len(trimmed) + 3
            data_sheet.cell(
                row=note_row, column=1,
                value=(f"این برگه تا سقف {MAX_EXPORT_ROWS:,} سطر صادر شده است؛ "
                       f"جدول کامل {len(frame):,} سطر دارد."))

    #: ---------------------------------------------------------- نوسان گروه‌ها
    if frame is not None and analysis.primary_metric and analysis.date_column \
            and analysis.primary_dimension:
        movement = movers(frame, analysis.primary_metric, analysis.date_column,
                          analysis.primary_dimension)
        rows = ([["بیشترین رشد", item["label"], item["before_text"],
                  item["after_text"], _percent(item["change"])]
                 for item in movement.get("gainers", [])]
                + [["بیشترین کاهش", item["label"], item["before_text"],
                    item["after_text"], _percent(item["change"])]
                   for item in movement.get("losers", [])])
        if rows:
            mover_sheet = book.create_sheet(_safe_sheet_title("نوسان گروه‌ها"))
            mover_sheet.sheet_view.rightToLeft = True
            _table(mover_sheet, 2, ["نوع", "گروه", "نیمهٔ اول", "نیمهٔ دوم",
                                    "تغییر"], rows, theme)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _percent(ratio: float | None) -> str:
    """نسبت به متن درصدی با علامت — برای جدول اکسل."""
    if ratio is None:
        return "—"
    return f"{'-' if ratio < 0 else '+'}{abs(ratio) * 100:.1f}%"


def filename(extension: str = "xlsx") -> str:
    """نام فایل خروجی اکسل."""
    return f"{OUTPUT_BASENAMES['excel']}.{extension}"
